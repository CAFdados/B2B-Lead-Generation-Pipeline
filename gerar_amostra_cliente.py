import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("\n[*] A iniciar geração da amostra de extração (Google Maps -> Emails)...")

# Dados de amostra altamente realistas (Agências em Dublin, cidade do cliente)
dados_amostra = [
    {"Nome da Empresa": "Dublin Digital Marketing", "Email": "hello@dublindigital.ie", "Telefone": "+353 1 555 0101", "Site": "https://dublindigital.ie", "Endereço": "12 O'Connell St Lower, Dublin", "Link Google Maps": "https://maps.google.com/?cid=123456789"},
    {"Nome da Empresa": "Celtic SEO Solutions", "Email": "contact@celticseo.com", "Telefone": "+353 1 555 0202", "Site": "https://celticseo.com", "Endereço": "45 Grafton Street, Dublin", "Link Google Maps": "https://maps.google.com/?cid=987654321"},
    {"Nome da Empresa": "Irish B2B Leads", "Email": "sales@irishb2b.ie", "Telefone": "+353 1 555 0303", "Site": "https://irishb2b.ie", "Endereço": "88 St Stephen's Green, Dublin", "Link Google Maps": "https://maps.google.com/?cid=456123789"},
    {"Nome da Empresa": "Growth Hackers DUB", "Email": "info@growthhackers.ie", "Telefone": "+353 1 555 0404", "Site": "https://growthhackers.ie", "Endereço": "Temple Bar, Dublin", "Link Google Maps": "https://maps.google.com/?cid=321654987"},
    {"Nome da Empresa": "Liffey Creative Agency", "Email": "hello@liffeycreative.com", "Telefone": "+353 1 555 0505", "Site": "https://liffeycreative.com", "Endereço": "1 Grand Canal Square, Dublin", "Link Google Maps": "https://maps.google.com/?cid=741852963"}
]

nome_arquivo = "Google_Maps_Email_Sample.xlsx"
df = pd.DataFrame(dados_amostra)
df.to_excel(nome_arquivo, index=False, engine='openpyxl')

# Aplicar o Design Premium que os clientes adoram
wb = openpyxl.load_workbook(nome_arquivo)
ws = wb.active
ws.title = "Google Maps Extraction"

header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid") # Azul corporativo profundo
header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
zebra_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    column_letter = get_column_letter(col)
    ws.column_dimensions[column_letter].width = 35 if col in [2, 4, 6] else 25 # Emails e links mais largos

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", horizontal="left")
        if cell.row % 2 == 0:
            cell.fill = zebra_fill

ws.freeze_panes = "A2"
wb.save(nome_arquivo)
print(f"[✔] Sucesso! A amostra '{nome_arquivo}' foi gerada e está pronta para ser enviada.\n")